# -*- coding: utf-8 -*-
"""Add or refresh the TTS model comparison block in 'AI_Uretim' (rows 35+). Re-runnable."""
from copy import copy
from pathlib import Path

import openpyxl

PATH = Path(__file__).resolve().parent.parent / "gbsoft-fiyat-modeli.xlsx"
SHEET = "AI_Uretim"
TOP = 35

# (model, USD / 1M karakter veya None, USD/dk doğrudan veya None, not)
MODELS = [
    ("OpenAI gpt-4o-mini-tts (sohbet — öneri)", None, 0.015,
     "Token bazlı; ~$0.015/dk üçüncü taraf kaynak — resmi sayfadan teyit et"),
    ("Azure Neural TTS", 15, None, "tr-TR sesleri var; üçüncü taraf kaynak — teyit et"),
    ("Google Chirp 3 HD", 30, None, "cloud.google.com/text-to-speech/pricing — ilk 1M karakter/ay ücretsiz"),
    ("ElevenLabs Flash/Turbo", 40, None, "elevenlabs.io/pricing/api — $0.04/1K; abonelik paketi düşürür"),
    ("ElevenLabs v3 / Multilingual v2 (anlatım — öneri)", 80, None,
     "elevenlabs.io/pricing/api — $0.08/1K; sabit anlatım için"),
]


def style(dst, src):
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.number_format = src.number_format
    dst.alignment = copy(src.alignment)


def main() -> None:
    wb = openpyxl.load_workbook(PATH)
    ws = wb[SHEET]
    sec, label, inp, note = ws["A4"], ws["A5"], ws["B5"], ws["C5"]
    hdr, usd, tot_l, tot_v = ws["A19"], ws["B27"], ws["A24"], ws["B24"]

    for row in ws.iter_rows(min_row=TOP, max_row=max(ws.max_row, TOP + 20)):
        for c in row:
            c.value = None

    def put(coord, value, like):
        ws[coord].value = value
        style(ws[coord], like)

    r = TOP
    put(f"A{r}", "TTS MODEL KARŞILAŞTIRMASI (fiyat kontrolü: Ekim 2026)", sec)
    put(f"A{r+1}", "Dakika başına karakter (tahmin)", label)
    put(f"B{r+1}", 1000, inp)
    put(f"C{r+1}", "Türkçe konuşma ≈ 900–1.000 karakter/dk", note)
    put(f"A{r+2}", "Sesi açan misafir oranı", label)
    put(f"B{r+2}", 1, inp)
    ws[f"B{r+2}"].number_format = "0%"
    put(f"C{r+2}", "1 = tavan senaryo; gerçekçi tahmin için 0,3 dene", note)
    chars, share = f"$B${r+1}", f"$B${r+2}"

    h = r + 4
    for col, text in zip("ABCDEFG", [
        "Model", "USD / 1M karakter", "USD / dakika", "Tur başı (USD)",
        "Otel başı / ay (USD)", "Platform / ay (USD)", "Not / kaynak",
    ]):
        put(f"{col}{h}", text, hdr)

    for i, (name, per_m, per_min, src) in enumerate(MODELS, start=h + 1):
        put(f"A{i}", name, label)
        if per_m is None:
            put(f"C{i}", per_min, inp)
            ws[f"C{i}"].number_format = "$0.000"
        else:
            put(f"B{i}", per_m, inp)
            ws[f"B{i}"].number_format = "$#,##0"
            put(f"C{i}", f"=B{i}*{chars}/1000000", usd)
            ws[f"C{i}"].number_format = "$0.000"
        put(f"D{i}", f"=C{i}*$B$8*{share}", usd)
        ws[f"D{i}"].number_format = "$0.000"
        put(f"E{i}", f"=D{i}*$B$15", usd)
        put(f"F{i}", f"=E{i}*Ozet!$B$8", usd)
        put(f"G{i}", src, note)

    n = h + len(MODELS) + 2
    put(f"A{n}", "ADIM ANLATIMI ÖN ÜRETİM (tek seferlik, otel başı)", sec)
    put(f"A{n+1}", "Otel başı anlatım karakteri (tüm adımlar × diller)", label)
    put(f"B{n+1}", 10000, inp)
    put(f"C{n+1}", "Tahmin: ~10 adım × 500 karakter × 2 dil; bir kez üretilip R2'de saklanır", note)
    put(f"A{n+2}", "Tek seferlik maliyet (ElevenLabs v3, USD)", tot_l)
    put(f"B{n+2}", f"=B{n+1}/1000000*B{h+len(MODELS)}", tot_v)

    g = n + 4
    put(f"A{g}", "KENDİ GPU ALTERNATİFİ (API yerine)", sec)
    put(f"A{g+1}", "Hetzner GEX44 (RTX 4000 Ada) — EUR/ay, GPU başı", label)
    put(f"B{g+1}", 184, inp)
    ws[f"B{g+1}"].number_format = "€#,##0"
    put(f"C{g+1}", "+ €79 kurulum. Model: Chatterbox Multilingual (MIT, Türkçe var)", note)
    put(f"A{g+2}", "RunPod RTX 4090 — USD/saat", label)
    put(f"B{g+2}", 0.6, inp)
    ws[f"B{g+2}"].number_format = "$0.00"
    put(f"C{g+2}", "Sürekli açık ≈ $440/ay; fiyat üçüncü taraf kaynak", note)
    put(f"A{g+3}", "Not", label)
    put(f"C{g+3}", "GPU sayısı ve Türkçe kalite ölçülmedi; ~40.000 dk/ay altında API daha ucuz (tahmin)", note)

    ws["C10"].value = "gpt-4o-mini-tts ≈ $0.015/dk (tahmini); alternatifler için aşağıdaki TTS tablosu"
    for col, width in {"C": 14, "D": 16, "F": 18, "G": 60}.items():
        if (ws.column_dimensions[col].width or 0) < width:
            ws.column_dimensions[col].width = width

    wb.save(PATH)
    print(f"Güncellendi: {PATH} → {SHEET} satır {TOP}+")


if __name__ == "__main__":
    main()
